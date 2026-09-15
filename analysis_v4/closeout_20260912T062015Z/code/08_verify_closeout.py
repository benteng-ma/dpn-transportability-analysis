from __future__ import annotations

import os
import ast,hashlib,json,re,subprocess
from pathlib import Path
import pandas as pd
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader

P = Path(os.environ.get("DPN_PROJECT_ROOT", Path(__file__).resolve().parents[3]))
T=P/"analysis_v4/targeted_repair_2026-09-12"; C=P/"analysis_v4/closeout_20260912T062015Z"; V=C/"09_verification"
tests=[]
def check(name,ok,detail): tests.append({"test":name,"passed":bool(ok),"detail":str(detail)})
def sha(p):
 h=hashlib.sha256()
 with open(p,"rb") as f:
  for b in iter(lambda:f.read(1048576),b""): h.update(b)
 return h.hexdigest()

for f in sorted((C/"code").glob("*.py")):
 try: ast.parse(f.read_text(encoding="utf-8")); check("python_ast_"+f.name,True,"parsed")
 except Exception as e: check("python_ast_"+f.name,False,e)

flow=pd.read_csv(C/"01_M04_flow/helper_flow_validation_retry/COUNT_FLOW.tsv",sep="\t")
fv=dict(zip(flow.stage,flow["count"])); check("M04_raw_partition",fv["unique_mapping"]+fv["ambiguous_or_unmapped"]==fv["all"],fv)
check("M04_eligible_geneids",fv["eligible_rank"]==27506,fv["eligible_rank"])
cache=pd.read_csv(C/"01_M04_flow/ANNOTATION_CACHE_VS_FIXED_SNAPSHOT.tsv",sep="\t")
check("M04_annotation_cache_exact",set(cache._merge)=={"both"},cache._merge.value_counts().to_dict())

dmr=pd.read_csv(C/"02_M02_audit/DMR_exact_reproduction_v2/DMR_REPRODUCTION_AUDIT.tsv",sep="\t")
check("DMR_coordinates_exact",dmr.coordinates_exact.all(),dmr.to_dict("records")); check("DMR_numeric_tolerance",dmr.max_numeric_abs_difference.max()<1e-12,dmr.max_numeric_abs_difference.max())
fx=pd.read_csv(C/"03_M02_fixed_regions/PROPENG_retest_v2/FIXED_REGION_RETEST_ALL.tsv",sep="\t")
check("fixed_region_family_167",len(fx)==167,len(fx)); check("fixed_region_all_evaluable",(fx.status=="COMPLETED").all(),fx.status.value_counts().to_dict()); check("fixed_region_no_q10",int((fx.BH_full_discovery_family<.1).sum())==0,fx.BH_full_discovery_family.min())

sp=pd.read_csv(C/"04_M03_spatial/SPATIAL_DONOR_EFFECTS.tsv",sep="\t")
elig=sp.eligible.astype(str).str.lower().eq("true"); check("spatial_algebra",((sp.loc[elig,"up_effect"]-sp.loc[elig,"down_effect"]-sp.loc[elig,"effect"]).abs()<1e-12).all(),"effect=up-down")
check("spatial_donors",sp.loc[elig,"donor_id"].nunique()==6,sp.loc[elig,"donor_id"].nunique()); check("spatial_P9_NE",((sp.program=="P9")&~elig).sum()==6,"six donor placeholders")
coords=pd.read_csv(C/"04_M03_spatial/M03_REAL_COORDINATES_WITH_AUTHOR_ROI.tsv.gz",sep="\t"); check("spatial_16_sections",coords.geo_accession.nunique()==16,coords.geo_accession.nunique())
con=pd.read_csv(C/"04_M03_spatial/M03_SECTION_GENE_CONTRIBUTIONS.tsv.gz",sep="\t",usecols=["geo_accession"]); check("spatial_gene_contributions_nonempty",len(con)==129537,len(con))

blocks=pd.read_csv(C/"05_block_summary/BLOCK_INFLUENCE_STANDARDIZED.tsv",sep="\t")
check("block_632",len(blocks)==632,len(blocks)); check("block_nine_fixed_colors",set(blocks.block)=={"black","blue","brown","green","magenta","pink","red","turquoise","yellow"},sorted(set(blocks.block)))
check("block_rank_universe_unchanged",blocks.rank_universe_same.astype(str).str.lower().eq("true").all(),"verified from targeted repair score caches")

claims=pd.read_csv(C/"00_admin/CLAIM_EVIDENCE_REGISTER.tsv",sep="\t")
missing=[x for x in claims.result_file if not (C/x).exists()]; check("claim_files_exist",not missing,missing)

for name in ["DPN_Closeout_Author_Review_Manuscript.docx","DPN_Closeout_Combined_Supplementary_Information.docx","DPN_Closeout_Cover_Letter_Draft.docx"]:
 d=Document(C/"08_manuscript"/name); headers=[p.text for s in d.sections for h in (s.header,s.first_page_header,s.even_page_header) for p in h.paragraphs if p.text.strip()]
 check(name+"_no_header",not headers,headers)
 bad=[]
 for ti,t in enumerate(d.tables):
  b=t._tbl.tblPr.find(qn("w:tblBorders"))
  if b is not None:
   for tag in ("left","right","insideH","insideV"):
    e=b.find(qn("w:"+tag))
    if e is not None and e.get(qn("w:val")) not in ("nil","none"): bad.append((ti,tag,e.get(qn("w:val"))))
 check(name+"_three_line_table_container",not bad,bad)
 if "Manuscript" in name and "Supplementary" not in name:
  abstract=d.paragraphs[6].text; keywords=d.paragraphs[7].text.split(":",1)[1].split(";")
  check("manuscript_abstract_le_200",len(abstract.split())<=200,len(abstract.split())); check("manuscript_keywords_le_6",len(keywords)<=6,len(keywords)); check("manuscript_display_items",len(d.inline_shapes)+len(d.tables)==8,(len(d.inline_shapes),len(d.tables)))

for name,expected in [("DPN_Closeout_Author_Review_Manuscript.pdf",35),("DPN_Closeout_Combined_Supplementary_Information.pdf",55)]:
 n=len(PdfReader(str(C/"08_manuscript"/name)).pages); check(name+"_pages",n==expected,n)
n=len(PdfReader(str(C/"10_readout/DPN_v4_CLOSEOUT_REVIEW_REPORT.pdf")).pages); check("review_report_pages",n==5,n)

inputs=[]
for f in [P/"analysis_v3/gene_definition_repair_v1/strict_background_addendum/inputs/REPAIRED_MEMBERS.tsv",T/"01_inputs/M04_ENSEMBL_ROW_MAPPING_AUDIT.tsv.gz",T/"07_resource_recovery/downloads/GSE286347_MatrixBetaVal.csv.gz",T/"07_resource_recovery/downloads/GSE295206_RAW.tar"]:
 inputs.append({"path":str(f),"bytes":f.stat().st_size,"sha256":sha(f)})
pd.DataFrame(inputs).to_csv(V/"INPUT_HASHES.tsv",sep="\t",index=False)
pd.DataFrame(tests).to_csv(V/"CLOSEOUT_TESTS.tsv",sep="\t",index=False)
(V/"CLOSEOUT_TEST_SUMMARY.json").write_text(json.dumps({"tests":len(tests),"passed":sum(x["passed"] for x in tests),"failed":sum(not x["passed"] for x in tests),"all_passed":all(x["passed"] for x in tests)},indent=2),encoding="utf-8")
print(json.dumps({"tests":len(tests),"failed":[x for x in tests if not x["passed"]]},indent=2))
